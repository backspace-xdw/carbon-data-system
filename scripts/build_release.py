# -*- coding: utf-8 -*-
"""
打包软著登记材料：

  1. 前端源代码.txt   —— 把 web/ 下的源码拼接成一个文本文件
  2. 后端源代码.txt   —— 把 server/src + prisma schema/seed 拼接
  3. 采集表.xlsx      —— 源代码采集表
  4. 双碳实时数据采集系统操作手册.docx  —— 把 docs/操作手册.md 转 Word

直接 python3 scripts/build_release.py 运行即可。
输出到项目根目录。
"""

import os
import re
import sys
from pathlib import Path
from datetime import date

ROOT = Path(__file__).resolve().parent.parent
SYSTEM_NAME = "双碳实时数据采集系统"
VERSION = "V2.0"
AUTHOR_NOTE = ""  # 留空，软著登记表里手工填

# 软著最少 30 页 + 30 页 = 60 页，每页 50 行，约 3000 行就够了。
# 但我们这次把全部源码都打印出来，方便审核。
FRONTEND_FILES = [
    # 入口页面
    "web/index.html",
    "web/login.html",
    # 样式
    "web/assets/css/reset.css",
    "web/assets/css/tokens.css",
    "web/assets/css/layout.css",
    "web/assets/css/components.css",
    "web/assets/css/login.css",
    # 主入口
    "web/assets/js/main.js",
    # 核心
    "web/assets/js/core/api.js",
    "web/assets/js/core/router.js",
    "web/assets/js/core/store.js",
    "web/assets/js/core/realtime.js",
    "web/assets/js/core/toast.js",
    "web/assets/js/core/modal.js",
    "web/assets/js/core/db.js",
    "web/assets/js/core/utils.js",
    # 自定义元素
    "web/assets/js/components/iecsp-card.js",
    "web/assets/js/components/iecsp-chart.js",
    # Canvas 图表
    "web/assets/js/charts/base.js",
    "web/assets/js/charts/line.js",
    "web/assets/js/charts/bar.js",
    "web/assets/js/charts/donut.js",
    "web/assets/js/charts/gauge.js",
    # 业务页面
    "web/assets/js/pages/cockpit.js",
    "web/assets/js/pages/energy.js",
    "web/assets/js/pages/meters.js",
    "web/assets/js/pages/carbon.js",
    "web/assets/js/pages/quota.js",
    "web/assets/js/pages/risk.js",
    "web/assets/js/pages/report.js",
    "web/assets/js/pages/account.js",
]

BACKEND_FILES = [
    # 配置 & 入口
    "server/src/main.ts",
    "server/src/config/env.ts",
    # 基础设施
    "server/src/infrastructure/logging/logger.ts",
    "server/src/infrastructure/persistence/prisma.ts",
    "server/src/infrastructure/timeseries/influx.ts",
    "server/src/infrastructure/messaging/mqtt.ts",
    # 领域类型
    "server/src/domain/identity/types.ts",
    "server/src/domain/energy/types.ts",
    "server/src/domain/carbon/factors.ts",
    # 应用层
    "server/src/application/identity/auth-service.ts",
    "server/src/application/account/account-service.ts",
    "server/src/application/meter/meter-service.ts",
    "server/src/application/energy/energy-service.ts",
    "server/src/application/carbon/carbon-service.ts",
    "server/src/application/quota/quota-service.ts",
    "server/src/application/risk/risk-service.ts",
    "server/src/application/report/report-service.ts",
    "server/src/application/cockpit/cockpit-service.ts",
    # 接口层
    "server/src/interface/plugins/auth-plugin.ts",
    "server/src/interface/http/auth-routes.ts",
    "server/src/interface/http/cockpit-routes.ts",
    "server/src/interface/http/meter-routes.ts",
    "server/src/interface/http/energy-routes.ts",
    "server/src/interface/http/carbon-routes.ts",
    "server/src/interface/http/quota-routes.ts",
    "server/src/interface/http/risk-routes.ts",
    "server/src/interface/http/report-routes.ts",
    "server/src/interface/http/account-routes.ts",
    "server/src/interface/ws/realtime.ts",
    # 数据库 schema 与种子
    "server/prisma/schema.prisma",
    "server/prisma/seed.ts",
]


def merge_files(files, out_path):
    parts = []
    total_lines = 0
    file_stats = []  # (filename, lines)
    for rel in files:
        p = ROOT / rel
        if not p.exists():
            print(f"  WARN 文件不存在：{rel}")
            continue
        text = p.read_text(encoding="utf-8")
        lines = text.count("\n") + (0 if text.endswith("\n") else 1)
        file_stats.append((rel, lines))
        total_lines += lines
        sep = f"\n\n/* ========== 文件: {rel} ========== */\n\n"
        parts.append(sep + text)
    out_path.write_text("".join(parts).lstrip(), encoding="utf-8")
    return total_lines, file_stats


def count_lines(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return sum(1 for _ in f)
    except Exception:
        return 0


def write_xlsx(frontend_total, backend_total, fe_stats, be_stats, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "源代码采集表"

    thin = Side(border_style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(name="宋体", size=11, bold=True)
    normal = Font(name="宋体", size=11)
    title_font = Font(name="宋体", size=14, bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF7")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 标题
    ws.merge_cells("A1:E1")
    ws["A1"] = "源代码采集表"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 28

    # 软件基本信息（两列布局：标签 | 值，4 行）
    info_rows = [
        ("软件名称", SYSTEM_NAME, "版本号", VERSION),
        ("开发语言", "TypeScript / JavaScript / HTML / CSS", "开发完成日期", date.today().strftime("%Y-%m-%d")),
        ("源代码总行数", str(frontend_total + backend_total), "其中前端 / 后端", f"{frontend_total} / {backend_total}"),
        ("代码采集说明", "全部源代码已按文件列出，每个文件以分隔符开头，便于核对", "", ""),
    ]
    r = 3
    for row in info_rows:
        ws.cell(row=r, column=1, value=row[0]).font = bold
        ws.cell(row=r, column=2, value=row[1]).font = normal
        ws.cell(row=r, column=3, value=row[2]).font = bold
        ws.cell(row=r, column=4, value=row[3]).font = normal
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = left if c % 2 == 0 else center
            if c % 2 == 1:
                cell.fill = head_fill
        # 第四行说明跨列
        if row[2] == "":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        r += 1

    # 表格分隔
    r += 1

    # 文件清单表头
    ws.cell(row=r, column=1, value="序号").font = bold
    ws.cell(row=r, column=2, value="模块").font = bold
    ws.cell(row=r, column=3, value="文件路径").font = bold
    ws.cell(row=r, column=4, value="代码行数").font = bold
    ws.cell(row=r, column=5, value="备注").font = bold
    for c in range(1, 6):
        cell = ws.cell(row=r, column=c)
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
    r += 1

    idx = 1
    for label, stats in (("前端", fe_stats), ("后端", be_stats)):
        for path, lines in stats:
            ws.cell(row=r, column=1, value=idx).alignment = center
            ws.cell(row=r, column=2, value=label).alignment = center
            ws.cell(row=r, column=3, value=path).alignment = left
            ws.cell(row=r, column=4, value=lines).alignment = center
            note = _file_note(path)
            ws.cell(row=r, column=5, value=note).alignment = left
            for c in range(1, 6):
                ws.cell(row=r, column=c).border = border
                ws.cell(row=r, column=c).font = normal
            idx += 1
            r += 1

    # 列宽
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 56
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 36

    wb.save(out_path)


def _file_note(path):
    """根据文件路径生成一句中文备注。"""
    name = os.path.basename(path)
    rules = [
        ("main.ts", "服务端启动入口"),
        ("main.js", "前端入口与导航"),
        ("env.ts", "环境变量"),
        ("logger.ts", "日志"),
        ("prisma.ts", "数据库客户端"),
        ("influx.ts", "时序库读写"),
        ("mqtt.ts", "MQTT 接入"),
        ("auth-plugin.ts", "登录鉴权插件"),
        ("auth-service.ts", "登录服务"),
        ("auth-routes.ts", "登录接口"),
        ("cockpit-service.ts", "总览统计"),
        ("cockpit-routes.ts", "总览接口"),
        ("meter-service.ts", "计量点档案"),
        ("meter-routes.ts", "计量点接口"),
        ("energy-service.ts", "能源数据"),
        ("energy-routes.ts", "能源数据接口"),
        ("carbon-service.ts", "碳排放核算"),
        ("carbon-routes.ts", "碳排放接口"),
        ("quota-service.ts", "配额管理"),
        ("quota-routes.ts", "配额接口"),
        ("risk-service.ts", "报警规则"),
        ("risk-routes.ts", "报警接口"),
        ("report-service.ts", "报送处理"),
        ("report-routes.ts", "报送接口"),
        ("account-service.ts", "账户管理"),
        ("account-routes.ts", "账户接口"),
        ("realtime.ts", "WebSocket 实时通道"),
        ("schema.prisma", "数据库 schema"),
        ("seed.ts", "初始化数据"),
        ("factors.ts", "排放因子默认值"),
        ("types.ts", "类型定义"),
        ("login.html", "登录页"),
        ("index.html", "应用入口页"),
        ("api.js", "接口调用"),
        ("router.js", "前端路由"),
        ("store.js", "状态容器"),
        ("realtime.js", "WebSocket 客户端"),
        ("toast.js", "提示组件"),
        ("modal.js", "对话框"),
        ("db.js", "本地缓存"),
        ("utils.js", "工具函数"),
        ("iecsp-card.js", "卡片组件"),
        ("iecsp-chart.js", "图表组件"),
        ("base.js", "图表基类"),
        ("line.js", "折线图"),
        ("bar.js", "柱状图"),
        ("donut.js", "环形图"),
        ("gauge.js", "仪表盘"),
        ("cockpit.js", "总览页面"),
        ("energy.js", "实时数据页面"),
        ("meters.js", "计量点页面"),
        ("carbon.js", "碳核算页面"),
        ("quota.js", "配额页面"),
        ("risk.js", "报警页面"),
        ("report.js", "报送页面"),
        ("account.js", "用户与组织页面"),
        ("reset.css", "样式重置"),
        ("tokens.css", "设计变量"),
        ("layout.css", "布局"),
        ("components.css", "组件样式"),
        ("login.css", "登录页样式"),
    ]
    for k, v in rules:
        if name == k:
            return v
    return ""


def md_to_docx(md_path, out_path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    # 全局字体
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    try:
        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:
        pass

    # 页面边距
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    lines = text.splitlines()
    i = 0
    n = len(lines)
    in_code = False
    code_buf = []

    while i < n:
        line = lines[i]
        # 代码块
        if line.startswith("```"):
            if in_code:
                # 结束代码块
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                try:
                    from docx.oxml.ns import qn
                    run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                except Exception:
                    pass
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 一级标题
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.strip() == "":
            doc.add_paragraph("")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        else:
            # 处理粗体 **xxx**
            p = doc.add_paragraph()
            _add_inline(p, line)
        i += 1

    doc.save(out_path)


def _add_inline(paragraph, text):
    from docx.shared import Pt
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def main():
    print(f"项目根目录: {ROOT}")

    print("\n[1/4] 生成 前端源代码.txt ...")
    fe_total, fe_stats = merge_files(FRONTEND_FILES, ROOT / "前端源代码.txt")
    print(f"  共 {len(fe_stats)} 个文件 / {fe_total} 行")

    print("\n[2/4] 生成 后端源代码.txt ...")
    be_total, be_stats = merge_files(BACKEND_FILES, ROOT / "后端源代码.txt")
    print(f"  共 {len(be_stats)} 个文件 / {be_total} 行")

    print("\n[3/4] 生成 采集表.xlsx ...")
    write_xlsx(fe_total, be_total, fe_stats, be_stats, ROOT / "采集表.xlsx")
    print("  done")

    print("\n[4/4] 生成 双碳实时数据采集系统操作手册.docx ...")
    md_to_docx(ROOT / "docs" / "操作手册.md", ROOT / "双碳实时数据采集系统操作手册.docx")
    print("  done")

    print("\n全部完成。输出文件位于：")
    for f in ["前端源代码.txt", "后端源代码.txt", "采集表.xlsx", "双碳实时数据采集系统操作手册.docx"]:
        p = ROOT / f
        print(f"  {p}  ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
